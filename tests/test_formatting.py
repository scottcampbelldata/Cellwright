# SPDX-License-Identifier: EUPL-1.2
import sys
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch

import openpyxl
from openpyxl.styles import Font
from migration import formatting, runtime, vba
from migration.convert import basic

SOURCE = '''Sub FormatSheet()
With Worksheets("Data=FY").Range("A1:B2")
 .Font.Bold = True
 With .Font
  .Color = RGB(255, 0, 1)
 End With
 .NumberFormat = "0.00"
End With
End Sub
'''


class FormattingTests(unittest.TestCase):
    def test_basic_and_python_rewrites_with_nested_font_and_quoted_equals(self):
        for converter in (lambda source: basic(source, []), vba.to_python):
            result = converter(SOURCE)
            self.assertFalse(result["todos"])
            self.assertEqual([item["rule_id"] for item in result["rewrites"] if item["rule_id"] != "FORMAT-WITH"],
                             ["FORMAT-FONT-BOLD", "FORMAT-FONT-COLOR", "FORMAT-NUMBERFORMAT"])
            self.assertIn("Data=FY", result["source"])
        self.assertIn('MFFormat mfwith2, "font.color", MFVbaRGB(255, 0, 1)', basic(SOURCE, [])["source"])
        self.assertIn("_vb.format_range(_with_2, 'font.color', _vb.rgb(255, 0, 1))", vba.to_python(SOURCE)["source"])
        compile(vba.to_python(SOURCE)["source"], "<formatting candidate>", "exec")

    def test_manual_properties_values_shadowing_and_unrelated_text(self):
        for statement in ('.Font.ColorIndex = 3', '.NumberFormat = "[$USD]0.00"', '.NumberFormatLocal = "0,00"', '.VerticalAlignment = 42', '.Interior.Color = -1', '.Font.Bold = Worksheets("Other").Range("B1").Font.Bold'):
            source = 'Sub Test()\nWorksheets("Data").Range("A1")' + statement + '\nEnd Sub'
            for converter in (lambda source: basic(source, []), vba.to_python):
                self.assertTrue(converter(source)["todos"], statement)
        unrelated = 'Sub Test()\nMsgBox "font and interior"\n\x27 .Font.Bold = True\nEnd Sub'
        self.assertFalse(basic(unrelated, [])["todos"])
        self.assertFalse(basic(unrelated, [])["rewrites"])
        value_with = 'Sub Test()\nWith Worksheets("Data").Range("A1")\n.Value = 1\nEnd With\nMsgBox ".Font.Bold"\nEnd Sub'
        self.assertIn(value_with, basic(value_with, [])["source"])
        self.assertFalse(basic(value_with, [])["todos"])
        compound = 'Sub Test()\nWorksheets("Data").Range("A1").Font.Bold = True: x = 1\nEnd Sub'
        self.assertTrue(basic(compound, [], native=True)["todos"])
        getter_with = SOURCE.replace('.Font.Bold = True', 'x = .Value')
        self.assertTrue(basic(getter_with, [])["todos"])
        shadowed = 'Const xlLeft = 7\n' + SOURCE
        self.assertTrue(basic(shadowed, [])["todos"])
        self.assertFalse(any(item['rule_id'].startswith('FORMAT-') for item in basic(shadowed, [])["rewrites"]))
        self.assertNotIn('MFVbaRGB', basic('Sub Test()\nWorksheets("Data").Range("A1").Font.Color = custom.RGB(1,2,3)\nEnd Sub', [])["rewrites"][0]["after"])

    def test_openpyxl_range_styles_colour_order_and_preserved_values(self):
        native = openpyxl.Workbook()
        native.active["A1"] = "kept"
        native.active["B2"] = "=1+2"
        native.active["A1"].font = Font(name="Existing", underline="single")
        cells = runtime.Workbook(native).worksheets(1).range("A1:B2")
        values = {"font.bold": True, "font.italic": -1, "font.name": "Liberation Sans", "font.size": 12,
                  "font.color": formatting.rgb(255, 0, 1), "interior.color": formatting.rgb(1, 0, 255),
                  "wraptext": True, "horizontalalignment": -4108, "verticalalignment": -4160, "numberformat": "0.00"}
        for prop, value in values.items(): runtime.format_range(cells, prop, value)
        for row in native.active["A1:B2"]:
            for cell in row:
                self.assertEqual(cell.font.color.rgb, "FFFF0001")
                self.assertEqual(cell.fill.fgColor.rgb, "FF0100FF")
                self.assertTrue(cell.font.bold and cell.font.italic and cell.alignment.wrap_text)
                self.assertEqual((cell.font.name, cell.font.size, cell.alignment.horizontal, cell.alignment.vertical, cell.number_format),
                                 ("Liberation Sans", 12, "center", "top", "0.00"))
        self.assertEqual(native.active["A1"].font.underline, "single")
        self.assertEqual(native.active["A1"].value, "kept")
        self.assertEqual(native.active["B2"].value, "=1+2")
        runtime.format_range(cells, "font.bold", False)
        runtime.format_range(cells, "font.italic", 0)
        runtime.format_range(cells, "wraptext", False)
        for row in native.active["A1:B2"]:
            for cell in row: self.assertFalse(cell.font.bold or cell.font.italic or cell.alignment.wrap_text)

    def test_invalid_values_and_cell_budget_fail_before_mutation(self):
        native = openpyxl.Workbook()
        sheet = runtime.Workbook(native).worksheets(1)
        for prop, value in (("font.color", -1), ("font.color", True), ("font.size", float('nan')), ("font.bold", "yes"), ("numberformat", "mm/dd/yyyy")):
            with self.assertRaises(NotImplementedError): runtime.format_range(sheet.range("A1"), prop, value)
        with self.assertRaises(ValueError): runtime.format_range(sheet.range("A1:A100001"), "font.bold", True)
        with self.assertRaises(NotImplementedError): formatting.rgb(256, 0, 0)
        self.assertFalse(native.active._cells)

    def test_uno_properties_enums_and_number_format_locale(self):
        cell = Mock()
        native_sheet = Mock(getCellRangeByPosition=Mock(return_value=cell))
        formats = Mock(queryKey=Mock(return_value=-1), addNew=Mock(return_value=73))
        book = SimpleNamespace(uno=True, native=SimpleNamespace(NumberFormats=formats))
        target = runtime.Range(SimpleNamespace(book=book, native=native_sheet), (1, 1, 2, 2))
        uno = SimpleNamespace(Enum=lambda kind, name: (kind, name), createUnoStruct=lambda kind: SimpleNamespace())
        with patch.dict(sys.modules, {"uno": uno}):
            runtime.format_range(target, "font.color", 255)
            runtime.format_range(target, "font.bold", True)
            runtime.format_range(target, "font.italic", True)
            runtime.format_range(target, "horizontalalignment", -4131)
            runtime.format_range(target, "verticalalignment", -4107)
            runtime.format_range(target, "numberformat", "0.00")
        self.assertEqual([call.args for call in cell.setPropertyValue.call_args_list], [
            ("CharColor", 0xFF0000), ("CharWeight", 150.0), ("CharPosture", ("com.sun.star.awt.FontSlant", "ITALIC")),
            ("HoriJustify", ("com.sun.star.table.CellHoriJustify", "LEFT")), ("VertJustify", 3), ("NumberFormat", 73)])
        locale = formats.queryKey.call_args.args[1]
        self.assertEqual((locale.Language, locale.Country), ("en", "US"))
