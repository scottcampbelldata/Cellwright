# SPDX-License-Identifier: EUPL-1.2
"""Generate and assess 30 owned business scenarios; this does not calibrate human effort."""
import argparse
import hashlib
import json
import math
import os
import shutil
import subprocess
import time
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

from scripts.generate_source_fixture import assemble
from migration.convert import convert
from migration import runtime

ROOT = Path(__file__).resolve().parents[1]
SAFE = {
    'budget': ('Dim i As Long\nFor i = 1 To 6\nbase = base + amount * i\nNext i', lambda base, amount: base + amount * 21),
    'payroll': ('base = base * amount * 0.8', lambda base, amount: base * amount * 0.8),
    'inventory': ('base = base - amount\nIf base < 0 Then\nbase = 0\nEnd If', lambda base, amount: max(0, base - amount)),
    'forecast': ('base = base * (1 + amount / 100)', lambda base, amount: base * (1 + amount / 100)),
    'invoice': ('base = base + amount * 0.2', lambda base, amount: base + amount * 0.2),
}
MANUAL = {
    'outlook-report': ('Sub Export()\nDim app As Object\nSet app = CreateObject("Outlook.Application")\nEnd Sub', {'MF002', 'MF004'}),
    'windows-api': ('Declare PtrSafe Sub Sleep Lib "kernel32" (ByVal milliseconds As Long)\nSub Wait()\nSleep 50\nEnd Sub', {'MF001'}),
    'database-refresh': ('Sub Refresh()\nDim db As Object\nSet db = CreateObject("ADODB.Connection")\nThisWorkbook.RefreshAll\nEnd Sub', {'MF002', 'MF008'}),
    'open-event': ('Sub Calculate()\nWorksheets("Data").Range("A1").Value = 100\nEnd Sub', set()),
}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--binary', type=Path, default=ROOT / ('target/release/macroferry.exe' if os.name == 'nt' else 'target/release/macroferry'))
    parser.add_argument('--generator', type=Path, default=ROOT / ('target/debug/examples/generate_vba_fixture.exe' if os.name == 'nt' else 'target/debug/examples/generate_vba_fixture'))
    parser.add_argument('--allow-execution', action='store_true', help='Run only the 15 owned arithmetic candidates through openpyxl')
    args = parser.parse_args()
    if not args.binary.is_file() or not args.generator.is_file(): parser.error('Build the scanner and generate_vba_fixture example first')
    output = args.out.resolve()
    output.mkdir(parents=True, exist_ok=False)
    books, sources = output / 'workbooks', output / 'sources'
    books.mkdir(); sources.mkdir()
    records = []
    for scenario in (*SAFE, *MANUAL, 'power-query'):
        for profile, rows in enumerate((12, 120, 1200)):
            name = f'{scenario}-{profile + 1}'
            if scenario == 'power-query':
                path = books / (name + '.xlsx')
                # The constant sales query is shared intentionally to exercise duplicated pipelines.
                shutil.copyfile(ROOT / 'corpus/generated/sales-query.xlsx', path)
                records.append({'file': path.name, 'scenario': scenario, 'profile': 'shared sales query', 'expected_rules': ['MF053']})
                continue
            book = openpyxl.Workbook(); data = book.active; data.title = 'Data'
            base, amount = 100 + profile * 20, (1.25, 0, -2)[profile]
            if scenario == 'inventory' and profile == 2: amount = 200  # Exercise a stock-out without producing negative inventory.
            data['B1'] = base
            transactions = book.create_sheet('Transactions')
            transactions.append(['Account', 'Units', 'Unit price', 'Amount'])
            for row in range(2, rows + 2):
                transactions.append([f'SIM-{row % 7:03}', row % 13 + 1, 12.5 + row % 11, f'=B{row}*C{row}'])
            table = Table(displayName='Transactions', ref=f'A1:D{rows + 1}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            transactions.add_table(table)
            summary = book.create_sheet('Summary'); summary['A1'] = 'Simulated business data'
            summary['B1'] = f'=SUM(Transactions!D2:D{rows + 1})'
            summary['A2'] = scenario; summary['A3'] = 'No production records'
            temporary = sources / (name + '.xlsx'); book.save(temporary); book.close()
            document = ''
            if scenario in SAFE:
                source = f'Option Explicit\nSub Calculate(ByVal amount As Double)\nDim base As Double\nbase = Worksheets("Data").Range("B1").Value\n{SAFE[scenario][0]}\nWorksheets("Data").Range("A1").Value = base\nWorksheets("Data").Range("C1").Formula = "=A1*2"\nEnd Sub\n'
                rules = set()
            else:
                source, rules = MANUAL[scenario]
                if scenario == 'open-event': document = 'Private Sub Workbook_Open()\nThisWorkbook.RefreshAll\nEnd Sub\n'
            source_path = sources / (name + '.bas'); source_path.write_text(source, encoding='utf-8')
            document_path = sources / (name + '-document.bas'); document_path.write_text(document, encoding='utf-8')
            binary = sources / (name + '.bin')
            subprocess.run([str(args.generator.resolve()), str(binary), str(source_path), str(document_path)], check=True, cwd=ROOT)
            path = books / (name + '.xlsm'); assemble(temporary, binary, path)
            records.append({'file': path.name, 'scenario': scenario, 'rows': rows, 'input': amount, 'base': base, 'expected_rules': sorted(rules)})
    started = time.monotonic()
    subprocess.run([str(args.binary.resolve()), 'scan', '--no-source', '--out', str(output / 'inventory'), '--', str(books)], check=True)
    scan = json.loads((output / 'inventory/scan.json').read_text(encoding='utf-8'))
    assert len(scan['workbooks']) == 30 and all(book['status'] == 'inspected' for book in scan['workbooks'])
    indexed = {Path(book['path']).name: book for book in scan['workbooks']}
    for record in records:
        book = indexed[record['file']]
        rules = {finding['rule_id'] for finding in book['findings']}
        rules.update(finding['rule_id'] for module in book['modules'] for finding in module['findings'])
        assert set(record['expected_rules']) <= rules, (record['file'], rules)
        if record['scenario'] == 'open-event': assert any('workbook_open' in entry.lower() for entry in book['entry_points']), book['entry_points']
        if record['scenario'] == 'power-query': assert book['queries'], record['file']
        record['verdict'] = book['verdict']; record['rules_found'] = sorted(rules)
        path = books / record['file']; before = hashlib.sha256(path.read_bytes()).hexdigest()
        destination = output / 'candidates' / path.stem
        candidate = convert(path, destination, args.binary, 'python')
        record['todos'] = sum(len(module['todos']) for module in candidate['modules']) + sum(len(query['todos']) for query in candidate.get('power_queries', []))
        record['execution'] = 'not run'
        if record['scenario'] in SAFE:
            module = next(module for module in candidate['modules'] if module['name'] == 'Arithmetic')
            assert not module['todos'], module['todos']
            if args.allow_execution:
                native = openpyxl.load_workbook(path, keep_vba=True)
                environment = runtime.Workbook(native).environment()
                exec((destination / module['candidate']).read_text(encoding='utf-8'), environment)
                environment['calculate'](record['input'])
                assert math.isclose(native['Data']['A1'].value, SAFE[record['scenario']][1](record['base'], record['input']), abs_tol=1e-9)
                assert native['Data']['C1'].value == '=A1*2'
                native.close(); record['execution'] = 'value and formula text passed'
        elif record['scenario'] == 'open-event': assert any('Events' in limitation for limitation in candidate['limitations'])
        elif record['scenario'] != 'power-query': assert record['todos'] > 0, record['file']
        record['manual_review_required'] = record['scenario'] in MANUAL
        assert hashlib.sha256(path.read_bytes()).hexdigest() == before, record['file']
    assert scan['duplicates'], 'Shared module templates should be deduplicated'
    summary = {'schema_version': '1', 'simulated': True, 'calibrated': False, 'measured_effort_hours': None,
               'workbooks': 30, 'converted': len(records), 'executed': sum(record['execution'] != 'not run' for record in records),
               'duplicate_groups': len(scan['duplicates']), 'scenario_checks_passed': True,
               'scope': 'Owned source-only VBA, business tables/formulas, COM/API blockers, document events and constant-record M. Formula recalculation and production effort are not measured.',
               'assessment_seconds': round(time.monotonic() - started, 3), 'cases': records}
    (output / 'summary.json').write_text(json.dumps(summary, indent=2) + '\n', encoding='utf-8')
    print(f"30 simulated workbooks assessed and converted; {summary['executed']} safe Python candidates checked. Estimates remain uncalibrated.")


if __name__ == '__main__': main()
