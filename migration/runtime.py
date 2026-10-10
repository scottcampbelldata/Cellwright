# SPDX-License-Identifier: EUPL-1.2
"""Small Excel-like adapter shared by generated UNO and openpyxl macros.

Unsupported operations raise; VBA coercion, error handling and UI are not emulated.
"""
import builtins
import datetime
import math
import re
from migration import formatting

rgb = formatting.rgb


def concat(left, right):
    return cstr(left) + cstr(right)


def idiv(left, right):
    return math.trunc(round(left) / round(right))


def mod(left, right):
    return round(left) - math.trunc(round(left) / round(right)) * round(right)


def for_range(start, end, step=1):
    if step == 0:
        raise ValueError("VBA For step cannot be zero")
    while start <= end if step > 0 else start >= end:
        yield start
        start += step


def cstr(value):
    if isinstance(value,float) and value.is_integer(): return str(builtins.int(value))
    return "" if value is None else str(value)


def and_(left,right): return left & right
def or_(left,right): return left | right
def not_(value): return not value if isinstance(value,bool) else ~value


def cint(value):
    result = int(round(float(value)))
    if not -32768 <= result <= 32767:
        raise OverflowError("VBA Integer overflow")
    return result


def clng(value):
    result = int(round(float(value)))
    if not -2147483648 <= result <= 2147483647:
        raise OverflowError("VBA Long overflow")
    return result


def cdbl(value):
    return float(value)


def val(value):
    match = re.match(r"\s*([+-]?(?:\d+(?:\.\d*)?|\.\d+))", str(value))
    return float(match[1]) if match else 0


def isnumeric(value):
    try:
        float(value)
        return not isinstance(value, bool)
    except (ValueError, TypeError):
        return False


def left(value, length):
    if length < 0:
        raise ValueError("Negative string length")
    return value[:length]


def right(value, length):
    if length < 0:
        raise ValueError("Negative string length")
    return value[-length:] if length else ""


def mid(value, start, length=None):
    if start < 1 or (length is not None and length < 0):
        raise ValueError("Invalid Mid arguments")
    return value[start - 1:] if length is None else value[start - 1:start - 1 + length]


def replace(value, find, replacement):
    return value.replace(find, replacement)


def instr(value, search):
    return value.find(search) + 1


def iif(condition, yes, no):
    return yes if condition else no


def array(*values):
    return list(values)


def lbound(value):
    return 0


def ubound(value):
    return len(value) - 1


def split(value, delimiter=" "):
    return [] if value == "" else value.split(delimiter)


def join(value, delimiter=" "):
    return delimiter.join(value)


def year(value): return value.year
def month(value): return value.month
def day(value): return value.day
def date(): return datetime.date.today()
def now(): return datetime.datetime.now()
def weekday(value): return (value.weekday() + 1) % 7 + 1


def dateadd(interval, amount, value):
    days = {"d": 1, "y": 1, "w": 1, "ww": 7, "h": 1 / 24, "n": 1 / 1440, "s": 1 / 86400}
    if interval not in days:
        raise NotImplementedError("DateAdd month/year intervals require review")
    return value + datetime.timedelta(days=round(amount) * days[interval])


abs = builtins.abs
len = builtins.len
round = builtins.round
int = math.floor
fix = math.trunc
sqr = math.sqrt
chr = builtins.chr
asc = lambda value: ord(value[0])
ucase = lambda value: value.upper()
lcase = lambda value: value.lower()
trim = lambda value: value.strip(" ")
ltrim = lambda value: value.lstrip(" ")
rtrim = lambda value: value.rstrip(" ")


def coordinates(address):
    match = re.fullmatch(r"\$?([A-Za-z]+)\$?(\d+)(?::\$?([A-Za-z]+)\$?(\d+))?", address)
    if not match:
        raise NotImplementedError("Only explicit A1 cell/range addresses are supported")

    def column(value):
        result = 0
        for char in value.upper():
            result = result * 26 + ord(char) - ord("A") + 1
        return result

    c1, r1 = column(match[1]), builtins.int(match[2])
    c2, r2 = column(match[3] or match[1]), builtins.int(match[4] or match[2])
    if min(r1, c1) < 1 or r2 < r1 or c2 < c1 or r2 > 1048576 or c2 > 16384:
        raise ValueError("Invalid Excel cell range")
    return r1, c1, r2, c2


def column_name(number):
    result = ""
    while number:
        number, digit = divmod(number - 1, 26)
        result = builtins.chr(65 + digit) + result
    return result


class Range:
    def __init__(self, sheet, bounds):
        self.sheet, self.bounds = sheet, bounds

    @property
    def count(self):
        r1, c1, r2, c2 = self.bounds
        return (r2 - r1 + 1) * (c2 - c1 + 1)

    @property
    def row(self): return self.bounds[0]

    @property
    def column(self): return self.bounds[1]

    @property
    def address(self):
        r1, c1, r2, c2 = self.bounds
        first, last = f"${column_name(c1)}${r1}", f"${column_name(c2)}${r2}"
        return first if first == last else first + ":" + last

    def _get(self, formula=False):
        r1, c1, r2, c2 = self.bounds
        rows = [[self.sheet.read(r, c, formula) for c in builtins.range(c1, c2 + 1)] for r in builtins.range(r1, r2 + 1)]
        return rows[0][0] if self.count == 1 else tuple(tuple(row) for row in rows)

    def _set(self, value, formula=False):
        r1, c1, r2, c2 = self.bounds
        sequence = isinstance(value, (tuple, list))
        if sequence and (len(value) != r2 - r1 + 1 or any(not isinstance(row, (tuple, list)) or len(row) != c2 - c1 + 1 for row in value)):
            raise ValueError("Range assignment dimensions differ")
        for r in builtins.range(r1, r2 + 1):
            for c in builtins.range(c1, c2 + 1):
                self.sheet.write(r, c, value[r - r1][c - c1] if sequence else value, formula)

    value = property(_get, _set)
    value2 = property(_get, _set)
    formula = property(lambda self: self._get(True), lambda self, value: self._set(value, True))

    def _bounded(self):
        if self.count > 100000: raise ValueError("Range exceeds 100,000 cells")
        r1,c1,r2,c2=self.bounds
        if min(r1,c1)<1 or r2>1048576 or c2>16384:raise ValueError("Range exceeds sheet bounds")
        return self.sheet.native.getCellRangeByName(self.address) if self.sheet.book.uno else None

    def clearcontents(self):
        native=self._bounded()
        if native is not None:native.clearContents(23)
        else:self._set(None)

    def clear(self):
        native=self._bounded()
        if native is not None:native.clearContents(887)
        else:
            for cell in self:
                r,c,_,_=cell.bounds
                target=self.sheet.native.cell(r,c)
                target.value=None;target._style=None;target.comment=None;target.hyperlink=None

    def sort(self, key1, order1, header):
        native=self._bounded()
        if not isinstance(key1,Range) or key1.sheet.name!=self.sheet.name or key1.sheet.book is not self.sheet.book:raise ValueError("Sort key must be in this worksheet")
        r1,c1,r2,c2=self.bounds
        _,key,_,end=key1.bounds
        if key!=end or not c1<=key<=c2 or type(order1) is not builtins.int or type(header) is not builtins.int or order1 not in (1,2) or header not in (1,2):raise ValueError("Explicit one-column key, order and header required")
        start=r1+(header==1)
        keys=[self.sheet.read(r,key) for r in builtins.range(start,r2+1)]
        if any(type(value) not in (builtins.int,float) or not math.isfinite(value) or isinstance(self.sheet.read(r,key,True),str) and self.sheet.read(r,key,True).startswith('=') for r,value in zip(builtins.range(start,r2+1),keys)):
            raise NotImplementedError("Only numeric constant sort keys supported")
        if native is not None:
            import uno
            field=uno.createUnoStruct('com.sun.star.table.TableSortField')
            field.Field=key-c1;field.IsAscending=order1==1
            field.FieldType=uno.Enum('com.sun.star.table.TableSortFieldType','NUMERIC')
            from harness.libreoffice.runner import property_value
            native.sort(tuple(property_value(name,value) for name,value in [('SortFields',uno.Any('[]com.sun.star.table.TableSortField',(field,))),('ContainsHeader',header==1),('IsSortColumns',False)]))
        else:
            from copy import copy
            rows=[]
            for row in builtins.range(start,r2+1):
                cells=[self.sheet.native.cell(row,col) for col in builtins.range(c1,c2+1)]
                if any(cell.data_type=='f' for cell in cells):raise NotImplementedError("Formula sorting requires an engine-native migration")
                rows.append((keys[row-start],[(cell.value,copy(cell._style),copy(cell.comment),copy(cell.hyperlink)) for cell in cells]))
            for row,(_,cells) in zip(builtins.range(start,r2+1),sorted(rows,key=lambda value:value[0],reverse=order1==2)):
                for col,(value,style,comment,link) in zip(builtins.range(c1,c2+1),cells):
                    cell=self.sheet.native.cell(row,col);cell.value=value;cell._style=style;cell.comment=comment;cell.hyperlink=link

    def autofilter(self, field, criteria1):
        native=self._bounded()
        r1,c1,r2,c2=self.bounds
        if type(field) is not builtins.int or not 1<=field<=c2-c1+1:raise ValueError("Invalid filter field")
        if type(criteria1) not in (str,builtins.int,float) or isinstance(criteria1,str) and (any(ch in criteria1 for ch in '*?~') or criteria1.startswith(('<','>','='))):
            raise NotImplementedError("Only literal equality filters supported")
        if native is not None:
            import uno
            condition=uno.createUnoStruct('com.sun.star.sheet.TableFilterField')
            condition.Field=field-1;condition.Operator=uno.Enum('com.sun.star.sheet.FilterOperator','EQUAL')
            condition.IsNumeric=type(criteria1) is not str
            if condition.IsNumeric:condition.NumericValue=criteria1
            else:condition.StringValue=criteria1
            descriptor=native.createFilterDescriptor(True)
            descriptor.ContainsHeader=True;descriptor.IsCaseSensitive=False;descriptor.UseRegularExpressions=False
            descriptor.setFilterFields((condition,));native.filter(descriptor)
        else:
            for row in builtins.range(r1+1,r2+1):
                value=self.sheet.read(row,c1+field-1)
                formula=self.sheet.read(row,c1+field-1,True)
                if isinstance(formula,str) and formula.startswith('='):raise NotImplementedError("Filtering formula results requires a calculating engine")
            self.sheet.native.auto_filter.ref=self.address
            for row in builtins.range(r1+1,r2+1):
                value=self.sheet.read(row,c1+field-1)
                equal=value.casefold()==criteria1.casefold() if isinstance(value,str) and isinstance(criteria1,str) else value==criteria1
                self.sheet.native.row_dimensions[row].hidden=not equal

    def offset(self, rows=0, columns=0):
        r1, c1, r2, c2 = self.bounds
        bounds = r1 + rows, c1 + columns, r2 + rows, c2 + columns
        if min(bounds) < 1:
            raise ValueError("Range offset before worksheet")
        return Range(self.sheet, bounds)

    def resize(self, rows, columns):
        if rows < 1 or columns < 1:
            raise ValueError("Range dimensions must be positive")
        r, c, _, _ = self.bounds
        return Range(self.sheet, (r, c, r + rows - 1, c + columns - 1))

    def copy(self, destination):
        native=self._bounded()
        if not isinstance(destination,Range) or destination.sheet.book is not self.sheet.book:raise ValueError("Copy destination must be in this workbook")
        r1, c1, r2, c2 = self.bounds
        destination.resize(r2-r1+1,c2-c1+1)._bounded()
        if native is not None:
            dest=destination.sheet.native.getCellRangeByName(destination.address).getCellByPosition(0,0).getCellAddress()
            destination.sheet.native.copyRange(dest,native.getRangeAddress())
            return
        values=self._get(True)
        cells=[values] if self.count==1 else [cell for row in values for cell in row]
        if any(isinstance(cell,str) and cell.startswith("=") for cell in cells):
            # shortcut: formula copies need relative-reference translation, use an engine-native implementation before enabling them.
            raise NotImplementedError("Formula copy and formatting require an engine-native migration")
        destination.resize(r2 - r1 + 1, c2 - c1 + 1)._set(values, True)

    def __iter__(self):
        r1, c1, r2, c2 = self.bounds
        for r in builtins.range(r1, r2 + 1):
            for c in builtins.range(c1, c2 + 1):
                yield Range(self.sheet, (r, c, r, c))


class Sheet:
    def __init__(self, book, native):
        self.book, self.native = book, native

    @property
    def name(self): return self.native.Name if self.book.uno else self.native.title

    @name.setter
    def name(self, value):
        if not isinstance(value,str) or not value or len(value)>31 or any(ch in value for ch in "[]:*?/\\") or value.startswith("'") or value.endswith("'"):
            raise ValueError("Invalid worksheet name")
        names=self.book.native.Sheets.getElementNames() if self.book.uno else self.book.native.sheetnames
        if any(name.casefold()==value.casefold() and name!=self.name for name in names):raise ValueError("Worksheet name already exists")
        if self.book.uno:self.native.Name=value
        else:self.native.title=value

    def range(self, address): return Range(self, coordinates(address))

    def cells(self, row, column):
        if row < 1 or column < 1:
            raise ValueError("Cell coordinates are one-based")
        return Range(self, (row, column, row, column))

    def read(self, row, column, formula=False):
        if self.book.uno:
            cell = self.native.getCellByPosition(column - 1, row - 1)
            if formula:
                return cell.Formula
            typ = str(cell.getType().value)
            return None if typ == "EMPTY" else cell.String if typ == "TEXT" else cell.Value
        cell = self.native.cell(row, column)
        if not formula and cell.data_type == "f":
            raise NotImplementedError("openpyxl cannot calculate formulas; use UNO for formula results")
        return cell.value

    def write(self, row, column, value, formula=False):
        if self.book.uno:
            cell = self.native.getCellByPosition(column - 1, row - 1)
            if formula and isinstance(value, str) and value.startswith("="):
                cell.Formula = value
            elif value is None:
                cell.clearContents(23)
            elif isinstance(value, (builtins.int, float, bool)):
                cell.Value = float(value)
            elif isinstance(value, str):
                cell.String = value
            else:
                raise NotImplementedError("UNO date/object range assignment requires review")
        else:
            self.native.cell(row, column).value = value

    def activate(self):
        self.book.active = self
        if self.book.uno:
            self.book.native.CurrentController.setActiveSheet(self.native)


def format_range(target, prop, value):
    if not isinstance(target, Range):
        raise NotImplementedError("Formatting requires a workbook range")
    value = formatting.validate(prop, value)
    r1, c1, r2, c2 = target.bounds
    if target.count > 100_000:
        raise ValueError("Formatting exceeds 100,000 cells")
    if target.sheet.book.uno:
        native = target.sheet.native.getCellRangeByPosition(c1 - 1, r1 - 1, c2 - 1, r2 - 1)
        if prop == "font.bold": value = 150.0 if value else 100.0
        elif prop == "font.italic":
            import uno
            value = uno.Enum("com.sun.star.awt.FontSlant", "ITALIC" if value else "NONE")
        elif prop == "horizontalalignment":
            import uno
            value = uno.Enum("com.sun.star.table.CellHoriJustify", formatting.HORIZONTAL[value][0])
        elif prop == "verticalalignment": value = formatting.VERTICAL[value][0]
        elif prop == "numberformat":
            import uno
            locale = uno.createUnoStruct("com.sun.star.lang.Locale")
            locale.Language, locale.Country = "en", "US"
            formats = target.sheet.book.native.NumberFormats
            key = formats.queryKey(value, locale, True)
            value = formats.addNew(value, locale) if key < 0 else key
        native.setPropertyValue(formatting.PROPERTIES[prop], value)
    else:
        from copy import copy
        from openpyxl.styles import PatternFill
        for row in target.sheet.native.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
            for cell in row:
                if prop.startswith("font."):
                    font = copy(cell.font)
                    setattr(font, {"font.bold": "bold", "font.italic": "italic", "font.name": "name", "font.size": "size", "font.color": "color"}[prop], f"FF{value:06X}" if prop.endswith("color") else value)
                    cell.font = font
                elif prop == "interior.color": cell.fill = PatternFill("solid", fgColor=f"FF{value:06X}")
                elif prop == "numberformat": cell.number_format = value
                else:
                    alignment = copy(cell.alignment)
                    attr = {"wraptext": "wrap_text", "horizontalalignment": "horizontal", "verticalalignment": "vertical"}[prop]
                    setattr(alignment, attr, formatting.HORIZONTAL[value][1] if prop == "horizontalalignment" else formatting.VERTICAL[value][1] if prop == "verticalalignment" else value)
                    cell.alignment = alignment


class WorksheetFunction:
    @staticmethod
    def _values(values):
        for value in values:
            if isinstance(value, Range): value = value.value
            if isinstance(value, (tuple, list)):
                yield from WorksheetFunction._values(value)
            elif isinstance(value, (builtins.int, float)) and not isinstance(value, bool):
                yield value

    def sum(self, *values): return sum(self._values(values))
    def average(self, *values):
        values = list(self._values(values))
        return sum(values) / len(values)
    def min(self, *values): return min(self._values(values))
    def max(self, *values): return max(self._values(values))
    def count(self, *values): return len(list(self._values(values)))


class Sheets:
    def __init__(self,book):self.book=book
    def __call__(self,key):return self.book._worksheet(key)
    def add(self):
        native=self.book.native
        names=native.Sheets.getElementNames() if self.book.uno else native.sheetnames
        index=1
        while 'Sheet'+str(index) in names:index+=1
        name='Sheet'+str(index)
        if self.book.uno:
            native.Sheets.insertNewByName(name,native.Sheets.Count)
            sheet=self(name)
            native.CurrentController.setActiveSheet(sheet.native)
        else:sheet=Sheet(self.book,native.create_sheet(name))
        self.book.active=sheet
        return sheet


class Workbook:
    def __init__(self, native, uno=False):
        self.native, self.uno = native, uno
        self.worksheets=Sheets(self)
        self.active = self.worksheets(1)
        self.worksheetfunction = WorksheetFunction()

    def _worksheet(self, key):
        if self.uno:
            sheets = self.native.Sheets
            native = sheets.getByIndex(key - 1) if isinstance(key, builtins.int) else sheets.getByName(key)
        else:
            native = self.native.worksheets[key - 1] if isinstance(key, builtins.int) else self.native[key]
        return Sheet(self, native)

    def environment(self):
        return {"thisworkbook": self, "activeworkbook": self, "application": self, "worksheets": self.worksheets, "sheets": self.worksheets, "range": lambda address: self.active.range(address), "cells": lambda row, column: self.active.cells(row, column), "activesheet": self.active, "worksheetfunction": self.worksheetfunction}
